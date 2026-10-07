"""
Sincroniza médicos, profesionales, prácticas y consultorios CEHTA desde Excel.

  python manage.py sync_cehta_excel [--dry-run] [path]
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from estudios.models import TipoEstudioComplementario
from estudios.practicas_cehta import CEHTA_PRACTICA_NOMBRES, slug_practica
from medicos.models import Especialidad, Medico
from profesionales.models import Profesional
from turnos.models import Recurso

DEFAULT_XLSX = Path('data/cehta/PROFESIONALES_PRACTICAS_CONSULTORIOS_CEHTA.xlsx')

ESPECIALIDADES_AMBITO_COMPLETO = frozenset(
    {
        'RESIDENTE',
        'CARDIOLOGIA',
        'CARDIOLOGÍA',
        'CIRUGIA GENERAL',
        'CIRUGÍA GENERAL',
    }
)

MEDICOS_FALLBACK = [
    ('RAMIREZ NOELIA', 'RESIDENTE'),
    ('SALMISTA ADRIANA', 'NEFROLOGÍA'),
    ('VILLECCO MIGUEL', 'TRAUMATOLOGÍA'),
    ('WILLIAMS DAVID', 'MEDICINA FAMILIAR'),
    ('YAMUNI RODRIGUES JUAN FRANCISCO', 'CLÍNICA MÉDICA'),
    ('QUIROGA MAURICIO', 'UROLOGÍA'),
    ('JOFRE MANUEL', 'RESIDENTE'),
    ('BLAYA PATRICIO AXELL', 'NEUROLOGÍA'),
    ('LLOYD JONES GERALDINE', 'RESIDENTE'),
    ('POLINI ALVARO CARLOS', 'TRAUMATOLOGÍA'),
    ('PIZZI ADRIAN', 'INMUNOLOGÍA Y ALERGIA'),
    ('PERALTA ALEJANDRO', 'CARDIOLOGÍA'),
    ('ARNOLDS ROY', 'CIRUGÍA GENERAL'),
    ('GERVASSONI LAURA MELISA', 'DIABETOLOGÍA'),
    ('GIRAUDO MARTIN', 'CARDIOLOGÍA'),
    ('GIAVEDONI NORMA', 'MEDICINA GENERAL'),
    ('BARDI JOSÉ', 'DIAGNOSTICO POR IMÁGENES'),
    ('HERRERA MIRIAM', 'CARDIOLOGÍA'),
    ('INGARAMO ROBERTO', 'CARDIOLOGÍA'),
    ('KIGUEL FLORENCIA', 'GINECOLOGÍA'),
    ('LAGIOIA ALBERTO', 'CARDIOLOGÍA'),
    ('LOMBARDO EDUARDO', 'RESIDENTE'),
    ('MOYA CECILIA', 'GINECOLOGÍA Y OBSTETRICIA'),
    ('INGARAMO CAROLINA', 'CARDIOLOGÍA'),
    ('LOPEZ NOEMI ALEJANDRA', 'CIRUGÍA GENERAL'),
]

PROFESIONALES_FALLBACK = [
    ('AGUADO MARICEL', 'PSICOLOGÍA'),
    ('WEISSBERG PAMELA NAHIR', 'NUTRICION'),
    ('CENDRA SILVANA', 'NUTRICION'),
    ('FIORDELLI FIORELLA', 'NUTRICION'),
    ('MACEIRO JOAQUIN', 'KINESIOLOGÍA'),
    ('MARTIN NICOLAS', 'CONTROL DE MARCAPASOS'),
]

CONSULTORIOS_FALLBACK = [
    ('1', 'consultorio'),
    ('2', 'consultorio'),
    ('3', 'consultorio'),
    ('4', 'consultorio'),
    ('5', 'consultorio'),
    ('6', 'consultorio'),
    ('7', 'consultorio'),
    ('8', 'consultorio'),
    ('9', 'consultorio'),
    ('EC', 'sala de ecografía'),
    ('PEG', 'sala de ergometría'),
]


def _norm_key(s: str) -> str:
    nfkd = unicodedata.normalize('NFKD', (s or '').strip().upper())
    return ''.join(c for c in nfkd if not unicodedata.combining(c))


def _parse_apellido_nombre(profesional: str) -> tuple[str, str]:
    parts = [p for p in (profesional or '').strip().split() if p]
    if not parts:
        return '', ''
    if len(parts) == 1:
        return parts[0].title(), ''
    return parts[0].title(), ' '.join(parts[1:]).title()


def _matricula_sintetica(prefix: str, apellido: str, nombre: str) -> str:
    base = _norm_key(f'{apellido}_{nombre}')
    base = re.sub(r'[^A-Z0-9]+', '_', base).strip('_')[:40] or 'X'
    return f'{prefix}-{base}'[:50]


def _ambito_for_especialidad(nombre: str) -> str:
    key = _norm_key(nombre)
    if key in {_norm_key(x) for x in ESPECIALIDADES_AMBITO_COMPLETO}:
        return Medico.AMBITO_COMPLETO
    return Medico.AMBITO_AMBULATORIO


def _recurso_nombre_cehta(codigo: str, descripcion: str) -> tuple[str, str]:
    code = str(codigo).strip()
    if code.upper() == 'EC':
        return 'Sala de ecografía', Recurso.TipoRecurso.SALA_PROCEDIMIENTO
    if code.upper() == 'PEG':
        return 'Sala de ergometría', Recurso.TipoRecurso.SALA_PROCEDIMIENTO
    # números 1..9
    try:
        n = int(float(code))
        return f'Consultorio {n}', Recurso.TipoRecurso.CONSULTORIO
    except (TypeError, ValueError):
        desc = (descripcion or 'consultorio').strip()
        return f'CEHTA {code} — {desc}', Recurso.TipoRecurso.CONSULTORIO


class Command(BaseCommand):
    help = 'Sincroniza Excel CEHTA (médicos, profesionales, prácticas, consultorios).'

    def add_arguments(self, parser):
        parser.add_argument(
            'xlsx_path',
            nargs='?',
            default=str(DEFAULT_XLSX),
            help=f'Ruta al Excel (default: {DEFAULT_XLSX})',
        )
        parser.add_argument('--dry-run', action='store_true', help='No escribe en la BD.')

    def handle(self, *args, **options):
        path = Path(options['xlsx_path']).expanduser()
        dry = options['dry_run']
        medicos_rows, prof_rows, practicas, consultorios = self._load_data(path)

        self.stdout.write(f'Fuente: {path if path.exists() else "fallback embebido"}')
        self.stdout.write(
            f'Médicos={len(medicos_rows)} Profesionales={len(prof_rows)} '
            f'Prácticas={len(practicas)} Consultorios={len(consultorios)}'
            + (' [DRY-RUN]' if dry else '')
        )

        if dry:
            self._preview(medicos_rows, prof_rows, practicas, consultorios)
            return

        with transaction.atomic():
            stats = {
                'esp_created': 0,
                'med_created': 0,
                'med_updated': 0,
                'prof_created': 0,
                'prof_updated': 0,
                'tipo_created': 0,
                'tipo_updated': 0,
                'rec_created': 0,
                'rec_updated': 0,
                'rec_deactivated': 0,
            }
            self._sync_medicos(medicos_rows, stats)
            self._sync_profesionales(prof_rows, stats)
            self._sync_practicas(practicas, stats)
            self._sync_consultorios(consultorios, stats)

        for k, v in stats.items():
            self.stdout.write(f'  {k}: {v}')
        self.stdout.write(self.style.SUCCESS('Sync CEHTA completado.'))

    def _load_data(self, path: Path):
        if path.exists():
            try:
                return self._load_xlsx(path)
            except Exception as exc:
                self.stdout.write(self.style.WARNING(f'No se pudo leer Excel ({exc}); uso fallback.'))
        return (
            MEDICOS_FALLBACK,
            PROFESIONALES_FALLBACK,
            list(dict.fromkeys(CEHTA_PRACTICA_NOMBRES)),
            CONSULTORIOS_FALLBACK,
        )

    def _load_xlsx(self, path: Path):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError('openpyxl no instalado') from exc

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

        def sheet_rows(name: str):
            if name not in wb.sheetnames:
                return []
            ws = wb[name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            # skip header
            out = []
            for r in rows[1:]:
                if not r or all(c is None or str(c).strip() == '' for c in r):
                    continue
                out.append(r)
            return out

        medicos = []
        for r in sheet_rows('MEDICOS'):
            medicos.append((str(r[0]).strip(), str(r[1]).strip() if len(r) > 1 and r[1] else ''))

        profesionales = []
        for r in sheet_rows('PROFESIONALES'):
            profesionales.append((str(r[0]).strip(), str(r[1]).strip() if len(r) > 1 and r[1] else ''))

        practicas = []
        seen = set()
        for r in sheet_rows('PRACTICAS'):
            nombre = str(r[0]).strip()
            if not nombre or nombre.upper() == 'PRACTICAS':
                continue
            key = _norm_key(nombre)
            if key in seen:
                continue
            seen.add(key)
            practicas.append(nombre)

        consultorios = []
        for r in sheet_rows('CONSULTORIOS'):
            codigo = r[0]
            if codigo is None:
                continue
            desc = str(r[1]).strip() if len(r) > 1 and r[1] else ''
            consultorios.append((str(codigo).strip(), desc))

        wb.close()
        if not practicas:
            practicas = list(dict.fromkeys(CEHTA_PRACTICA_NOMBRES))
        return medicos, profesionales, practicas, consultorios

    def _preview(self, medicos, profesionales, practicas, consultorios):
        self.stdout.write('— Médicos (muestra) —')
        for row in medicos[:5]:
            self.stdout.write(f'  {row}')
        self.stdout.write('— Profesionales —')
        for row in profesionales:
            self.stdout.write(f'  {row}')
        self.stdout.write(f'— Prácticas: {len(practicas)} —')
        self.stdout.write('— Consultorios —')
        for row in consultorios:
            nombre, tipo = _recurso_nombre_cehta(row[0], row[1])
            self.stdout.write(f'  {nombre} ({tipo})')

    def _get_or_create_esp(self, nombre: str, stats: dict) -> Especialidad | None:
        nombre = (nombre or '').strip()
        if not nombre:
            return None
        existing = Especialidad.objects.filter(nombre__iexact=nombre).first()
        if existing:
            return existing
        esp = Especialidad.objects.create(
            nombre=nombre.title() if nombre == nombre.upper() else nombre,
            descripcion=f'Especialidad importada CEHTA — {nombre}',
        )
        stats['esp_created'] += 1
        return esp

    def _sync_medicos(self, rows, stats):
        for profesional, especialidad in rows:
            apellido, nombre = _parse_apellido_nombre(profesional)
            if not apellido:
                continue
            esp = self._get_or_create_esp(especialidad, stats)
            ambito = _ambito_for_especialidad(especialidad)
            med = (
                Medico.objects.filter(apellido__iexact=apellido, nombre__iexact=nombre)
                .order_by('id')
                .first()
            )
            if med:
                changed = False
                if esp and med.especialidad_id != esp.id:
                    med.especialidad = esp
                    changed = True
                if med.ambito_atencion != ambito:
                    med.ambito_atencion = ambito
                    changed = True
                if changed:
                    med.save()
                    stats['med_updated'] += 1
            else:
                matricula = _matricula_sintetica('CEHTA', apellido, nombre)
                if Medico.objects.filter(matricula=matricula).exists():
                    matricula = f'{matricula}-{Medico.objects.count() + 1}'[:50]
                Medico.objects.create(
                    apellido=apellido,
                    nombre=nombre,
                    matricula=matricula,
                    especialidad=esp,
                    ambito_atencion=ambito,
                )
                stats['med_created'] += 1

    def _sync_profesionales(self, rows, stats):
        for profesional, especialidad in rows:
            apellido, nombre = _parse_apellido_nombre(profesional)
            if not apellido:
                continue
            esp = self._get_or_create_esp(especialidad, stats)
            prof = (
                Profesional.objects.filter(apellido__iexact=apellido, nombre__iexact=nombre)
                .order_by('id')
                .first()
            )
            if prof:
                if esp and prof.especialidad_id != esp.id:
                    prof.especialidad = esp
                    prof.save(update_fields=['especialidad', 'ultima_actualizacion'])
                    stats['prof_updated'] += 1
            else:
                matricula = _matricula_sintetica('PROF', apellido, nombre)
                if Profesional.objects.filter(matricula=matricula).exists():
                    matricula = f'{matricula}-{Profesional.objects.count() + 1}'[:50]
                Profesional.objects.create(
                    apellido=apellido,
                    nombre=nombre,
                    matricula=matricula,
                    especialidad=esp,
                )
                stats['prof_created'] += 1

    def _sync_practicas(self, practicas, stats):
        for nombre in practicas:
            codigo = slug_practica(nombre)[:50]
            tipo, created = TipoEstudioComplementario.objects.update_or_create(
                codigo=codigo,
                defaults={
                    'nombre': nombre[:255],
                    'practica': codigo,
                    'requiere_informe': True,
                    'activo': True,
                    'descripcion': 'Práctica CEHTA',
                },
            )
            if created:
                stats['tipo_created'] += 1
            else:
                stats['tipo_updated'] += 1

    def _sync_consultorios(self, consultorios, stats):
        keep_names: set[str] = set()
        for codigo, desc in consultorios:
            nombre, tipo = _recurso_nombre_cehta(codigo, desc)
            keep_names.add(nombre)
            rec, created = Recurso.objects.update_or_create(
                nombre=nombre,
                defaults={
                    'ubicacion': Recurso.Ubicacion.CEHTA,
                    'tipo_recurso': tipo,
                    'activo': True,
                },
            )
            if created:
                stats['rec_created'] += 1
            else:
                stats['rec_updated'] += 1

        qs = Recurso.objects.filter(ubicacion=Recurso.Ubicacion.CEHTA).exclude(
            nombre__in=keep_names
        )
        deactivated = qs.filter(activo=True).update(activo=False)
        stats['rec_deactivated'] += deactivated
