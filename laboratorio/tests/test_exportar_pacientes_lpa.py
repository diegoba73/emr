"""Tests export Excel pacientes LPA (sin MedGemma)."""
from decimal import Decimal
from pathlib import Path
import tempfile

import pytest
from django.core.management import call_command
from django.test import TestCase

from laboratorio.management.commands.exportar_pacientes_lpa import (
    pacientes_con_lpa_ids,
)
from laboratorio.models import ResultadoExamen, SolicitudExamen, TipoExamen, TipoMuestra
from pacientes.models import Paciente


@pytest.mark.django_db
class TestExportPacientesLpa(TestCase):
    def setUp(self):
        self.tm = TipoMuestra.objects.create(codigo="SUERO_LPA", nombre="Suero", activo=True)
        self.te_lpa = TipoExamen.objects.create(
            codigo="LPA", nombre="Lipoproteína A", tipo_muestra_requerida=self.tm, precio=1, activo=True
        )
        self.te_glu = TipoExamen.objects.create(
            codigo="GLU", nombre="Glucosa", tipo_muestra_requerida=self.tm, precio=1, activo=True
        )
        self.pac_con = Paciente.objects.create(dni="30111222", nombre="Ana", apellido="Lpa")
        self.pac_sin = Paciente.objects.create(dni="30999888", nombre="Sin", apellido="Lpa")

        sol = SolicitudExamen.objects.create(
            paciente=self.pac_con,
            origen_solicitud="AMBULATORIO_CEHTA",
            estado="EN_PROCESO",
            numero="LAB-LPA-001",
        )
        ResultadoExamen.objects.create(
            solicitud=sol,
            tipo_examen=self.te_lpa,
            valor_obtenido="45",
            valor_numerico=Decimal("45"),
        )
        ResultadoExamen.objects.create(
            solicitud=sol,
            tipo_examen=self.te_glu,
            valor_obtenido="90",
            valor_numerico=Decimal("90"),
        )
        SolicitudExamen.objects.filter(pk=sol.pk).update(estado="FINALIZADO")

        sol2 = SolicitudExamen.objects.create(
            paciente=self.pac_sin,
            origen_solicitud="AMBULATORIO_CEHTA",
            estado="EN_PROCESO",
            numero="LAB-LPA-002",
        )
        ResultadoExamen.objects.create(
            solicitud=sol2,
            tipo_examen=self.te_glu,
            valor_obtenido="100",
            valor_numerico=Decimal("100"),
        )

    def test_filtra_solo_pacientes_con_lpa(self):
        ids = pacientes_con_lpa_ids()
        self.assertEqual(ids, {self.pac_con.id})

    def test_comando_escribe_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "lpa.xlsx"
            call_command("exportar_pacientes_lpa", output=str(out))
            self.assertTrue(out.is_file())
            import openpyxl

            wb = openpyxl.load_workbook(out, read_only=True)
            self.assertIn("todos_examenes", wb.sheetnames)
            self.assertIn("solo_lpa", wb.sheetnames)
            ws = wb["todos_examenes"]
            rows = list(ws.iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "dni")
            dnis = {r[0] for r in rows[1:]}
            self.assertEqual(dnis, {"30111222"})
            wb.close()
