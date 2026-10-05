"""
Cohortes por filtros de exámenes (Filtros avanzados).

modo ``all`` (Todos): pacientes con al menos un valor en cada código seleccionado.
modo ``any`` (Algunos):
  - sin obligatorios → al menos uno de los códigos (OR clásico);
  - con obligatorios → la cohorte son solo quienes tienen todos los obligatorios
    (AND). Los demás códigos seleccionados no filtran: van como columnas extra
    en el Excel y se completan si esa orden también los tiene.

Excel: una fila por orden que cumpla los códigos de filtro de fila
(todos los seleccionados en modo all; los obligatorios en modo any;
cualquiera de los códigos si any sin obligatorios). Columnas = todos los
códigos seleccionados (vacío si esa orden no tiene el examen).
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable, Literal

from django.db.models import Count
from django.utils import timezone

from laboratorio.models import ResultadoExamen, TipoExamen

ModoFiltro = Literal["all", "any"]


def _norm_codigos(codigos: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for raw in codigos:
        c = (raw or "").strip().upper()
        if not c or c in seen:
            continue
        seen.add(c)
        out.append(c)
    return out


def _fmt_fecha(dt: datetime | date | None) -> str:
    if dt is None:
        return ""
    if isinstance(dt, datetime):
        if timezone.is_aware(dt):
            dt = timezone.localtime(dt)
        return dt.strftime("%Y-%m-%d %H:%M")
    return dt.isoformat()


def _fecha_resultado(r: ResultadoExamen) -> datetime | None:
    return r.fecha_validacion or getattr(r.solicitud, "fecha_solicitud", None)


def resolver_examenes(codigos: list[str]) -> list[dict[str, Any]]:
    """Lista ordenada codigo+nombre (catálogo); códigos desconocidos se listan igual."""
    if not codigos:
        return []
    by_code = {
        (t.codigo or "").strip().upper(): (t.nombre or "").strip()
        for t in TipoExamen.objects.filter(codigo__in=codigos)
    }
    return [
        {"codigo": c, "nombre": by_code.get(c) or c}
        for c in codigos
    ]


def _qs_con_valor(codigos: list[str], *, desde: date | None, hasta: date | None):
    qs = ResultadoExamen.objects.exclude(valor_obtenido="").filter(
        tipo_examen__codigo__in=codigos
    )
    if desde is not None:
        qs = qs.filter(solicitud__fecha_solicitud__date__gte=desde)
    if hasta is not None:
        qs = qs.filter(solicitud__fecha_solicitud__date__lte=hasta)
    return qs


def _ids_con_todos(codigos: list[str], *, desde: date | None, hasta: date | None) -> set[int]:
    if not codigos:
        return set()
    cov = (
        _qs_con_valor(codigos, desde=desde, hasta=hasta)
        .values("solicitud__paciente_id")
        .annotate(n_codigos=Count("tipo_examen__codigo", distinct=True))
        .filter(n_codigos__gte=len(codigos))
    )
    return {row["solicitud__paciente_id"] for row in cov}


def _ids_con_alguno(codigos: list[str], *, desde: date | None, hasta: date | None) -> set[int]:
    if not codigos:
        return set()
    return set(
        _qs_con_valor(codigos, desde=desde, hasta=hasta)
        .values_list("solicitud__paciente_id", flat=True)
        .distinct()
    )


def paciente_ids_por_filtro(
    codigos: list[str],
    *,
    modo: ModoFiltro = "all",
    obligatorios: Iterable[str] | None = None,
    desde: date | None = None,
    hasta: date | None = None,
) -> set[int]:
    codigos = _norm_codigos(codigos)
    if not codigos:
        return set()

    if modo == "all":
        return _ids_con_todos(codigos, desde=desde, hasta=hasta)

    # Algunos: solo los obligatorios definen la cohorte; el resto es solo columna.
    req = _norm_codigos(obligatorios or [])
    req = [c for c in req if c in set(codigos)]
    if req:
        return _ids_con_todos(req, desde=desde, hasta=hasta)
    return _ids_con_alguno(codigos, desde=desde, hasta=hasta)


def resumen_filtro(
    codigos: list[str],
    *,
    modo: ModoFiltro = "all",
    obligatorios: Iterable[str] | None = None,
    desde: date | None = None,
    hasta: date | None = None,
) -> dict[str, Any]:
    codigos = _norm_codigos(codigos)
    obl = _norm_codigos(obligatorios or [])
    obl = [c for c in obl if c in set(codigos)]
    ids = paciente_ids_por_filtro(
        codigos, modo=modo, obligatorios=obl, desde=desde, hasta=hasta
    )
    examenes = resolver_examenes(codigos)
    for e in examenes:
        e["obligatorio"] = True if modo == "all" else e["codigo"] in set(obl)
    return {
        "modo": modo,
        "codigos": codigos,
        "obligatorios": obl if modo == "any" else list(codigos),
        "examenes": examenes,
        "desde": desde.isoformat() if desde else None,
        "hasta": hasta.isoformat() if hasta else None,
        "pacientes": len(ids),
        "marcado_estudio": True,
    }


def filas_ancho(
    codigos: list[str],
    paciente_ids: set[int],
    *,
    requeridos: Iterable[str] | None = None,
    desde: date | None = None,
    hasta: date | None = None,
) -> list[dict[str, Any]]:
    """
    Filas para Excel/JSON: una por solicitud del paciente.

    ``requeridos``: códigos que esa orden debe tener con valor para aparecer
    (p. ej. obligatorios en modo Algunos, o todos en modo Todos). Si vacío,
    basta con que la orden tenga ≥1 de ``codigos``.
    Las columnas son siempre todos los ``codigos`` (vacío si no está en la orden).
    """
    codigos = _norm_codigos(codigos)
    if not codigos or not paciente_ids:
        return []

    req = _norm_codigos(requeridos or [])
    req = [c for c in req if c in set(codigos)]

    qs = (
        ResultadoExamen.objects.filter(
            solicitud__paciente_id__in=paciente_ids,
            tipo_examen__codigo__in=codigos,
        )
        .exclude(valor_obtenido="")
        .select_related("solicitud", "solicitud__paciente", "tipo_examen")
        .order_by("solicitud__paciente_id", "solicitud_id", "tipo_examen__codigo")
    )
    if desde is not None:
        qs = qs.filter(solicitud__fecha_solicitud__date__gte=desde)
    if hasta is not None:
        qs = qs.filter(solicitud__fecha_solicitud__date__lte=hasta)

    buckets: dict[int, dict[str, Any]] = {}
    for r in qs.iterator(chunk_size=3000):
        sol = r.solicitud
        sid = sol.id
        if sid not in buckets:
            pac = sol.paciente
            buckets[sid] = {
                "dni": pac.dni or "",
                "apellido": pac.apellido or "",
                "nombre": pac.nombre or "",
                "fecha_analisis": _fmt_fecha(_fecha_resultado(r) or sol.fecha_solicitud),
                "protocolo": sol.numero or "",
                "estado_orden": sol.estado or "",
                "valores": {},
                "_sort": (
                    pac.dni or "",
                    sol.fecha_solicitud.isoformat() if sol.fecha_solicitud else "",
                    sol.numero or "",
                ),
            }
        codigo = (r.tipo_examen.codigo or "").strip().upper()
        if codigo and codigo not in buckets[sid]["valores"]:
            val = (r.valor_obtenido or "").strip()
            if r.unidad:
                val = f"{val} {r.unidad}".strip()
            buckets[sid]["valores"][codigo] = val

    rows: list[dict[str, Any]] = []
    for bucket in sorted(buckets.values(), key=lambda b: b["_sort"]):
        vals = bucket["valores"]
        if req:
            if any(c not in vals for c in req):
                continue
        elif not vals:
            continue
        row = {
            "dni": bucket["dni"],
            "apellido": bucket["apellido"],
            "nombre": bucket["nombre"],
            "fecha_analisis": bucket["fecha_analisis"],
            "protocolo": bucket["protocolo"],
            "estado_orden": bucket["estado_orden"],
        }
        for c in codigos:
            row[c] = vals.get(c, "")
        rows.append(row)
    return rows


def escribir_xlsx_bytes(
    codigos: list[str],
    rows: list[dict[str, Any]],
    *,
    meta: dict[str, Any],
) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Falta openpyxl") from exc

    codigos = _norm_codigos(codigos)
    examenes = resolver_examenes(codigos)
    headers = [
        "dni",
        "apellido",
        "nombre",
        "fecha_analisis",
        "protocolo",
        "estado_orden",
        *codigos,
    ]
    header_labels = [
        "DNI",
        "Apellido",
        "Nombre",
        "Fecha análisis",
        "Protocolo",
        "Estado orden",
        *[
            f"{e['codigo']}"
            + (f" — {e['nombre']}" if e["nombre"] != e["codigo"] else "")
            for e in examenes
        ],
    ]

    wb = openpyxl.Workbook()
    ws_meta = wb.active
    ws_meta.title = "resumen"
    ws_meta["A1"] = "Filtros avanzados — listado para estudio"
    ws_meta["A1"].font = Font(bold=True)
    r = 2
    for k, v in meta.items():
        ws_meta.cell(row=r, column=1, value=str(k))
        ws_meta.cell(row=r, column=2, value=str(v))
        r += 1

    ws = wb.create_sheet("listado")
    for col, label in enumerate(header_labels, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, start=2):
        for col, key in enumerate(headers, start=1):
            ws.cell(row=i, column=col, value=row.get(key, ""))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
